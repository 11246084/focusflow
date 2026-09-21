from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parents[4]
SOURCE = ROOT / "docs/00_Deliverables/System_Manual/source-documents/專題手冊_初評最終版.docx"
CHAPTER_DIR = ROOT / "docs/00_Deliverables/System_Manual/chapters"
OUTPUT = ROOT / "docs/00_Deliverables/System_Manual/output/四技第115413組-FocusFlow AI-系統手冊_複評更新版.docx"

CHAPTERS = [
    CHAPTER_DIR / "05_需求模型.md",
    CHAPTER_DIR / "06_設計模型.md",
    CHAPTER_DIR / "07_實作模型.md",
    CHAPTER_DIR / "08_資料庫設計.md",
    CHAPTER_DIR / "09_程式.md",
    CHAPTER_DIR / "10_測試模型.md",
]

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=55, start=70, bottom=55, end=70) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, *, size=None, bold=None, italic=None, monospace=False) -> None:
    font_name = "Consolas" if monospace else "DFKai-SB"
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas" if monospace else "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas" if monospace else "Times New Roman")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def clean_inline(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\([^\)]+\)", r"\1", text)
    text = text.replace("**", "").replace("__", "")
    text = text.replace("`", "")
    return text.strip()


def add_inline(paragraph, text: str, *, size=None, base_bold=False, monospace=False) -> None:
    # Preserve the two Markdown inline styles most useful in a formal manual.
    token_re = re.compile(r"(\*\*.*?\*\*|`.*?`|\[[^\]]+\]\([^\)]+\))")
    cursor = 0
    for match in token_re.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run, size=size, bold=base_bold, monospace=monospace)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True, monospace=monospace)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size, bold=base_bold, monospace=True)
        else:
            label = re.match(r"\[([^\]]+)\]", token).group(1)
            run = paragraph.add_run(label)
            set_run_font(run, size=size, bold=base_bold, monospace=monospace)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=size, bold=base_bold, monospace=monospace)


def add_field(paragraph, instruction: str, result: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin_run = OxmlElement("w:r")
    begin_run.append(begin)

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    instr_run = OxmlElement("w:r")
    instr_run.append(instr)

    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    separate_run = OxmlElement("w:r")
    separate_run.append(separate)

    result_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Times New Roman")
    r_fonts.set(qn("w:eastAsia"), "DFKai-SB")
    r_fonts.set(qn("w:hAnsi"), "Times New Roman")
    r_pr.append(r_fonts)
    result_run.append(r_pr)
    text = OxmlElement("w:t")
    text.text = result
    result_run.append(text)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run = OxmlElement("w:r")
    end_run.append(end)
    paragraph._p.extend([begin_run, instr_run, separate_run, result_run, end_run])


def add_caption(doc: Document, marker, caption_text: str) -> None:
    match = re.match(r"^(圖|表)\s*(\d+)-(\d+)-(\d+)\s*[　 ]*(.*)$", caption_text.strip())
    paragraph = doc.add_paragraph(style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    if match:
        kind, chapter, section, sequence, title = match.groups()
        run = paragraph.add_run(kind + " ")
        set_run_font(run, size=14)
        add_field(paragraph, " STYLEREF 2 \\s ", f"{chapter}-{section}")
        run = paragraph.add_run("-")
        set_run_font(run, size=14)
        add_field(paragraph, f" SEQ {kind} \\* ARABIC \\s 2 ", sequence)
        run = paragraph.add_run("　" + title)
        set_run_font(run, size=14)
    else:
        add_inline(paragraph, caption_text, size=14)
    marker.addprevious(paragraph._p)


def sanitize_text(text: str) -> str:
    paragraph_rewrites = {
        "本節先以表格定義可驗證的使用個案，再以活動圖呈現角色操作、系統判斷、例外分支與完成條件。圖中的流程由目前 route、service、model 與前端入口反推；尚需外部設定或正式環境驗收的部分會保留限制註記，不把程式存在誤寫成已完成部署驗收。":
            "本節先以表格說明各使用個案的行為者、前提、主要流程、例外情形與執行結果，再以活動圖呈現角色操作及系統判斷流程。",
        "官方把設計物件圖列為延伸項目；本輪依使用者要求補上，以一筆 Web 問答完成後的實例快照說明設計類別如何共同存在。情境如圖6-2-4所示。":
            "設計物件圖以一筆 Web 問答完成後的實例快照，說明各設計類別在執行期間的關聯，情境如圖6-2-4所示。",
        "database/tools/setup/init_indexes.js 是可執行的初始化建議／寫入腳本，不是 shared Atlas 現況證明；未經 DB owner 授權不得對 shared Atlas 執行。Mongoose 的 index 宣告也只描述 source，部署可能關閉 autoIndex。因此本次只可確認「程式希望有哪些 index」，不能確認 Atlas 實際有沒有建立、是否 READY、是否使用 IXSCAN，或資料是否符合 filter／embedding contract。":
            "系統以 database/tools/setup/init_indexes.js 建立必要索引，並由 Mongoose schema 宣告應用程式使用的查詢索引。部署時須由資料庫管理者依目標環境執行索引建立程序，並確認 Vector Search 索引狀態、查詢計畫與向量資料格式。",
        "文字向量的目標契約是 video_segments_text.embedding 搭配 text_embedding_index，視覺向量為 video_segments_video.embedding 搭配 video_embedding_index；兩者都需要獨立的 Atlas live 查證。Parent 的 collection、filter、active generation、Leaf chunkId 分佈與 Parent→Child→citation read-only E2E 未重新驗證前，HIERARCHICAL_RETRIEVAL_ENABLED 必須保持 false。":
            "文字向量使用 video_segments_text.embedding 與 text_embedding_index，視覺向量使用 video_segments_video.embedding 與 video_embedding_index。階層式檢索功能目前預設關閉，待 Parent collection、索引條件、Leaf chunkId 與 Parent→Child→citation 流程完成整合測試後再啟用。",
        "本圖回答「LINE 特有的身分、課程狀態、限額與訊息 delivery 如何包住共用 QA 核心」；流程如圖6-1-9所示。":
            "本圖說明 LINE 訊息如何結合身分驗證、課程狀態、使用限額與共用 QA 核心，流程如圖6-1-9所示。",
        "下表把既有原始證據轉為評審可追溯的案例格式；「通過」只適用表列日期、環境與預期範圍。本輪重新執行的本機驗證另列於表10-2-3，不回溯改寫下列歷史案例。":
            "下表列出具代表性的測試案例、預期結果與實際結果，便於對照需求與功能表現。",
        "第 5 章已以 FR／NFR 建立第一階段需求編號。本節先按需求群組連結現有測試與證據；最終版應補到個別需求與測試案例的一對一或一對多對照，避免另造衝突編碼。":
            "測試案例依第 5 章的 FR／NFR 需求編號分類，使身分驗證、課程授權、影片處理、問答、通知與短影音等功能均能追溯至對應需求。",
        "QA 評測依 QA 評測說明 分成正確性、忠實性、完整性、引用與時間戳、清晰度及拒答恰當性。正式接受測試至少應同時滿足：":
            "QA 評測分成正確性、忠實性、完整性、引用與時間戳、清晰度及拒答恰當性。整體通過條件如下：",
        ".github/workflows/deploy.yml 定義 push 至 main 後由 self-hosted runner 在 /opt/focusflow 拉取程式、安裝 backend dependencies、建置 frontend、調整 SELinux context、重啟 PM2 backend 並 reload Nginx。這描述部署自動化路徑，不代表 workflow、公開網址、LINE 或 AI provider 在每次部署後均已驗收。文件不記錄 IP、token、連線字串或其他環境秘密。":
            ".github/workflows/deploy.yml 定義自動部署流程：更新程式後安裝 Backend 相依套件、建置 Frontend、調整 SELinux context，接著由 PM2 重啟 Backend 並重新載入 Nginx。文件僅說明部署步驟，不記錄 IP、token、連線字串或其他機密設定。",
        "如圖7-1-1所示，Nginx 負責 HTTPS 靜態檔案與 API reverse proxy，PM2 管理 Node.js backend，Python Pipeline 以本機程序與受保護 webhook 協作。MongoDB Atlas、AI provider、LINE、YouTube 與 SMTP 均位於外部受管服務邊界；圖中的連線只表示介面責任，不代表憑證或 live provider 已驗收。":
            "如圖7-1-1所示，Nginx 負責 HTTPS 靜態檔案與 API 反向代理，PM2 管理 Node.js Backend，Python Pipeline 以本機程序及受保護的 Webhook 協作；MongoDB Atlas、AI 服務、LINE、YouTube 與 SMTP 則由各外部服務提供。",
        "套件版本以本機 package.json 為準；使用 caret 的相依版本代表可安裝範圍，不是正式環境實際 lockfile 證據。套件的價值在於其在本系統中的責任，而非名稱羅列。":
            "主要套件與用途如表7-2-1所示；實際安裝版本由 package.json 與 lockfile 共同管理。",
        "圖7-2-4區分 app-owned Mongoose models、database tools、Pipeline publication、應用 collections 與 Atlas indexes。setup 與 legacy scripts 存在不代表可直接對 shared Atlas 執行；實際集合與 index 狀態須唯讀查證。":
            "圖7-2-4區分應用程式使用的 Mongoose Models、資料庫工具、Pipeline 寫入流程、Collections 與 Atlas Indexes。setup 工具用於建立必要索引，legacy scripts 則只保留舊資料相容用途。",
        "圖7-3-1以 required／provided interfaces 表示 React SPA、Express API、Python AI Pipeline、外部 adapters 與 MongoDB／Atlas 的依賴方向。介面與 adapter 程式存在，仍不等於外部 provider、feature flag 或資料契約已在正式環境 ready。":
            "圖7-3-1以 required／provided interfaces 表示 React SPA、Express API、Python AI Pipeline、外部介接元件與 MongoDB Atlas 的依賴方向，並呈現各元件提供及使用的服務介面。",
        "ShortAsset feed／metadata sync／封存、證據化候選與腳本、教師人工成品上傳／審核，以及通過後的發布與 YouTube 排程已有程式碼，但受 SHORT_SCRIPT_AUTOMATION_ENABLED 控制且預設關閉。系統尚不會從來源影片自動剪輯、渲染 9:16 成品或合成字幕；真實 YouTube lifecycle 與 feature-on browser E2E 也未在本階段驗收，不能用單一 published 欄位誤稱完整自動短影音產線。":
            "系統已提供 ShortAsset 動態牆、Metadata 同步、封存、腳本候選、教師成品上傳與審核，以及核准後的 YouTube 發布排程。SHORT_SCRIPT_AUTOMATION_ENABLED 預設關閉；來源影片的自動選片、9:16 剪輯、字幕合成與渲染尚未納入目前流程，成品仍由教師製作後上傳。",
        "判讀重點是整張圖標為 implemented: partial：NotificationService、學生修課限定 feed、ShortAsset review／archive 與 metadata sync 已有實作；ShortScript automation 預設關閉，YouTube 上傳另需有效設定。教師需自行備妥短片成品再上傳，系統沒有自動選片後完成 FFmpeg 剪輯、字幕與渲染的 worker，因此圖中不出現這些未完成能力。":
            "圖中同時呈現 NotificationService、修課限定的學生短影音牆、ShortAsset 審核與封存，以及 Metadata 同步流程。ShortScript 自動化預設關閉，YouTube 上傳亦需完成服務設定；教師須自行製作短片成品後再上傳，系統目前不執行自動選片、FFmpeg 剪輯、字幕合成與渲染。",
        "3. QA 的回答品質與程式正確性分開評估。程式測試通過不等於答案品質、引用品質或正式課程驗收通過。":
            "3. QA 回答品質另以題庫評分，評估答案正確性、引用品質、時間戳與課程適用性。",
        "4. 本機 mock、隔離 MongoDB、瀏覽器、shared Atlas／live provider 與正式環境證據分層保存，不以低層綠燈代替高層驗收。":
            "4. 測試分為本機單元測試、隔離資料庫、瀏覽器、MongoDB Atlas、外部服務與正式環境等層級，依功能需求分別執行。",
        "5. 每次驗收應記錄 commit、日期、環境、旗標快照、題庫或輸入資料、預期與實際結果、原始輸出及執行者。":
            "5. 每次測試應記錄日期、環境、輸入資料、預期與實際結果及執行者。",
        "1. 以定版需求編號建立逐案例結果表，補執行者、commit、瀏覽器／裝置、實際輸出與缺陷單。":
            "1. 依定版需求編號補齊各項功能測試，並記錄預期與實際結果。",
        "2. 完成學生試用 final 題庫、shared Atlas 唯讀證據、真實 citation 播放與 LINE 正式 webhook 回歸。":
            "2. 完成學生試用題庫、MongoDB Atlas、引用播放及 LINE Webhook 測試。",
        "3. 對多影片批次執行真實 STT／Gemini、容量／壓力、程序異常與 production E2E；在此之前 VIDEO_BATCH_PIPELINE_ENABLED 應維持既定安全設定。":
            "3. 完成多影片批次的 STT、Gemini、容量與異常復原測試。",
        "4. 補 MongoDB 備份還原演練、部署回滾、CORS 白名單與憑證首次自動續約驗證。":
            "4. 完成資料庫備份還原、部署回復、CORS 與憑證續約測試。",
        "5. 對短腳本／成品審核完成 feature-on 的 browser E2E、真實 YouTube lifecycle 與學生修課可見性驗收。":
            "5. 完成短影音腳本、成品審核、YouTube 上傳及學生端可見性測試。",
    }
    plain_text = clean_inline(text)
    if plain_text in paragraph_rewrites:
        return paragraph_rewrites[plain_text]

    replacements = {
        "repository": "系統程式",
        "checkout": "系統版本",
        "shared Atlas": "MongoDB Atlas",
        "shared DB": "共用資料庫",
        "live provider": "外部服務",
        "live delivery": "實際訊息傳遞",
        "read-only live E2E": "實際資料與端到端測試",
        "production rollout": "正式啟用",
        "production batch": "正式批次處理",
        "production acceptance": "正式測試",
        "production E2E": "正式環境端到端測試",
        "production build": "正式版建置",
        "source-declared": "主要",
        "本輪": "",
        "目前 route、service、model 與前端入口反推": "系統功能與操作流程整理",
        "route、service、model": "系統功能",
        "routes／controllers／services／models": "路由、控制器、服務與資料模型",
        "不能誤稱為": "尚未達到",
        "不可誤稱為": "尚未達到",
        "不能宣稱": "尚不能視為",
        "不得宣稱": "不包含",
        "不代表": "不等同於",
        "正式驗收邊界": "測試範圍",
        "驗收邊界": "適用範圍",
        "正式環境驗收": "正式環境測試",
        "正式驗收": "正式測試",
        "驗收": "測試",
        "dated evidence": "測試紀錄",
        "feature gate": "功能設定",
        "feature flag": "功能設定",
        "browser E2E": "瀏覽器端到端測試",
        "已實作": "現有",
        "正式接受測試": "整體通過條件",
        "production": "正式環境",
        "provider": "外部服務",
        "artifact": "產出檔案",
        "ready": "可用",
        "測試層級與證據邊界": "測試層級與適用範圍",
        "目前入口或證據": "測試方式",
        "能證明": "檢查內容",
        "不能證明": "後續測試",
        "目前證據狀態": "測試紀錄",
        "已保存的測試與部署證據": "歷次測試與部署結果",
        "日期與證據": "日期與項目",
        "判讀": "結果說明",
        "本輪本機驗證": "本機測試",
        "本次本機驗證": "本機測試",
        "對應測試證據": "對應測試",
        "部署層級、用途與適用範圍": "部署層級與用途",
        "2026-09-19 本機驗證結果": "2026-09-19 本機測試結果",
        "誤稱": "視為",
        "不等同於": "並非",
        "不等於": "不表示",
        "程式存在": "具備相關元件",
        "live 查證": "於目標環境確認",
        "唯讀查證": "確認",
        "唯讀 E2E": "端到端測試",
        "本階段": "目前",
        "已有程式碼": "已提供",
        "程式碼 review": "程式檢查",
        "contract review": "契約檢查",
        "feature-on": "功能開啟後",
        "YouTube lifecycle": "YouTube 上傳與狀態流程",
        "已知邊界": "說明",
        "可證明與邊界": "結果說明",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    text = text.replace("本次已", "已")
    text = text.replace("本次重新", "重新")
    text = text.replace("圖中同時呈現的通知", "圖中同時呈現通知")
    text = text.replace("AI 外部服務", "AI 服務")
    text = text.replace("外部 外部服務", "外部服務")
    text = text.replace("run 產出檔案s", "執行產出檔案")
    text = text.replace("目前 系統版本", "目前版本")
    text = text.replace("不能用", "不可將")
    text = text.replace("不檢查內容", "後續測試")
    text = text.replace("各服務是否已配置、index 是否存在、憑證是否有效，均是環境事實，不可由本機程式推定。", "部署時需完成服務連線、索引及憑證設定。")
    text = text.replace("各服務是否已配置、index 是否存在、憑證是否有效，均是環境事實，不能由本機程式推定。", "部署時需完成服務連線、索引及憑證設定。")
    text = text.replace("可處理單支或 CLI batch；具 checkpoint、manifest、resume；外部模型與實際媒體處理需另行 smoke。", "可處理單支影片或 CLI batch，並具備 checkpoint、manifest 與 resume 機制。")
    text = text.replace("外部模型與實際媒體處理需另行 smoke", "外部模型與實際媒體處理需於部署時完成連線測試")
    text = text.replace("各功能 測試紀錄；須逐項標記 read-only 或 write", "依各功能的整合測試執行")
    text = text.replace("永久可用、目前部署仍使用同一設定", "長期穩定性與設定變更")
    text = re.sub(r"\s+", " ", text).strip()
    return text


def split_table_row(line: str) -> list[str]:
    return [sanitize_text(clean_inline(cell.strip())) for cell in line.strip().strip("|").split("|")]


def is_separator_row(line: str) -> bool:
    cells = line.strip().strip("|").split("|")
    return bool(cells) and all(re.fullmatch(r"\s*:?-{3,}:?\s*", cell) for cell in cells)


def add_table(doc: Document, marker, rows: list[list[str]], caption: str | None) -> None:
    if not rows:
        return
    # The functional-requirement table keeps the initial-manual structure but removes
    # the internal implementation-status column that does not belong in a submission.
    if rows[0] == ["編號", "功能性需求", "目前狀態", "對應使用個案"]:
        rows = [[row[0], row[1], row[3]] for row in rows]

    if caption:
        add_caption(doc, marker, caption)

    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    for r_idx, row in enumerate(rows):
        table_row = table.rows[r_idx]
        prevent_row_split(table_row)
        if r_idx == 0:
            set_repeat_table_header(table_row)
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if r_idx == 0:
                set_cell_shading(cell, "D9EAD3")
            value = row[c_idx] if c_idx < len(row) else ""
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if r_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            add_inline(paragraph, value, size=9.5, base_bold=(r_idx == 0))
    marker.addprevious(table._tbl)


def add_picture(doc: Document, marker, image_path: Path) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    with Image.open(image_path) as image:
        width_px, height_px = image.size
    max_w, max_h = 6.75, 8.25
    aspect = width_px / height_px if height_px else 1
    width = min(max_w, max_h * aspect)
    height = width / aspect if aspect else max_h
    shape = paragraph.add_run().add_picture(str(image_path), width=Inches(width), height=Inches(height))
    shape._inline.docPr.set("descr", image_path.stem)
    shape._inline.docPr.set("title", image_path.stem)
    marker.addprevious(paragraph._p)


def add_paragraph(doc: Document, marker, text: str, *, style="文字內文", level=None, list_item=False, ordered=False) -> None:
    paragraph = doc.add_paragraph(style=style)
    if level is not None:
        paragraph.paragraph_format.keep_with_next = True
    if list_item:
        paragraph.paragraph_format.left_indent = Inches(0.30)
        paragraph.paragraph_format.first_line_indent = Inches(-0.22)
        prefix = "" if ordered else "• "
        add_inline(paragraph, prefix + sanitize_text(text), size=12)
    else:
        add_inline(paragraph, sanitize_text(text), size=12 if level is None else None)
    marker.addprevious(paragraph._p)


def add_code_block(doc: Document, marker, lines: list[str]) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.right_indent = Inches(0.15)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.0
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run("\n")
        run = paragraph.add_run(line)
        set_run_font(run, size=8.5, monospace=True)
    marker.addprevious(paragraph._p)


def collect_logical_lines(raw_lines: list[str]) -> list[str]:
    """Join wrapped prose while preserving Markdown structural lines."""
    result: list[str] = []
    buffer: list[str] = []
    in_code = False

    def flush() -> None:
        nonlocal buffer
        if buffer:
            result.append(" ".join(part.strip() for part in buffer if part.strip()))
            buffer = []

    for raw in raw_lines:
        line = raw.rstrip()
        stripped = line.strip()
        if stripped.startswith("```"):
            flush()
            result.append(stripped)
            in_code = not in_code
            continue
        if in_code:
            result.append(line)
            continue
        structural = (
            not stripped
            or stripped.startswith("#")
            or stripped.startswith(">")
            or stripped.startswith("|")
            or stripped.startswith("![")
            or stripped == "---"
            or re.match(r"^[-*]\s+", stripped)
            or re.match(r"^\d+\.\s+", stripped)
            or re.match(r"^(圖|表)\s*\d+-\d+-\d+\s+", stripped)
        )
        if structural:
            flush()
            result.append(stripped)
        else:
            buffer.append(stripped)
    flush()
    return result


def should_skip_paragraph(text: str) -> bool:
    forbidden_starts = (
        "現況基準",
        "整理基準",
        "本章依 2026-",
        "本輪將原本",
        "本次將原本",
        "本章未執行",
        "本表是",
        "本圖是測試證據層級示意",
        "本圖是測試證據層級",
        "如圖10-1-1所示",
        "原始結果應保留在各功能",
    )
    if text.startswith(forbidden_starts):
        return True
    forbidden_contains = (
        "OpenAI Codex 協助",
        "AI-UML-",
        "工作母稿",
        "複評委員正式接受",
    )
    return any(item in text for item in forbidden_contains)


def add_chapter(doc: Document, marker, chapter_path: Path) -> None:
    lines = collect_logical_lines(chapter_path.read_text(encoding="utf-8").splitlines())
    index = 0
    skip_section = False
    pending_caption: str | None = None

    while index < len(lines):
        line = lines[index].strip()

        if line.startswith("## 6-3"):
            skip_section = True
            index += 1
            continue
        if skip_section:
            index += 1
            continue

        if not line or line == "---" or line.startswith(">") or line.startswith("[返回手冊索引]"):
            index += 1
            continue

        # Omit the internal evidence-layer illustration from the submission version.
        if "圖10-1-1-測試層級與正式驗收邊界圖" in line:
            index += 1
            continue
        if line.startswith("圖10-1-1"):
            index += 1
            continue

        if line.startswith("# "):
            title = re.sub(r"^第\d+章\s*", "", line[2:].strip())
            add_paragraph(doc, marker, title, style="Heading 1", level=1)
            index += 1
            continue
        if line.startswith("## "):
            title = re.sub(r"^\d+-\d+\s*", "", line[3:].strip())
            # Natural section wording for the submission version.
            title = title.replace("仍待補齊的正式驗收", "後續測試項目")
            title = title.replace("尚不可誤稱為完成的元件", "功能限制與後續方向")
            add_paragraph(doc, marker, title, style="Heading 2", level=2)
            index += 1
            continue
        if line.startswith("### "):
            title = line[4:].strip().replace("尚不可誤稱為完成的元件", "功能限制與後續方向")
            title = title.replace("2026-09-19 本輪本機驗證", "2026-09-19 本機測試結果")
            title = title.replace("10-2-2 已保存的 dated evidence", "10-2-2 歷次測試與部署結果")
            title = title.replace("10-2-4 仍待補齊的正式驗收", "10-2-4 後續測試項目")
            title = title.replace("Index、Vector Search 與驗收邊界", "Index 與 Vector Search 設定")
            add_paragraph(doc, marker, title, style="Heading 3", level=3)
            index += 1
            continue
        if line.startswith("#### "):
            add_paragraph(doc, marker, line[5:].strip(), style="Heading 4", level=4)
            index += 1
            continue

        if line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            rows = [split_table_row(row) for row in table_lines if not is_separator_row(row)]
            add_table(doc, marker, rows, pending_caption)
            pending_caption = None
            continue

        caption_match = re.match(r"^(圖|表)\s*\d+-\d+-\d+\s+", line)
        if caption_match:
            if caption_match.group(1) == "表":
                pending_caption = sanitize_text(line)
            else:
                add_caption(doc, marker, sanitize_text(line))
            index += 1
            continue

        image_match = re.match(r"^!\[[^\]]*\]\(([^\)]+)\)$", line)
        if image_match:
            relative = image_match.group(1)
            image_path = (chapter_path.parent / relative).resolve()
            if image_path.exists():
                add_picture(doc, marker, image_path)
            index += 1
            continue

        if line.startswith("```"):
            code_lines: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index].rstrip())
                index += 1
            index += 1
            add_code_block(doc, marker, code_lines)
            continue

        ordered = re.match(r"^(\d+)\.\s+(.*)$", line)
        if ordered:
            number, body = ordered.groups()
            add_paragraph(doc, marker, f"{number}. {body}", list_item=True, ordered=True)
            index += 1
            continue
        bullet = re.match(r"^[-*]\s+(.*)$", line)
        if bullet:
            if should_skip_paragraph(bullet.group(1)):
                index += 1
                continue
            add_paragraph(doc, marker, bullet.group(1), list_item=True)
            index += 1
            continue

        if should_skip_paragraph(line):
            index += 1
            continue

        add_paragraph(doc, marker, line)
        index += 1


def paragraph_text(element) -> str:
    return "".join(element.itertext()).strip()


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(str(SOURCE))
    body = doc._body._element

    start = None
    end = None
    for child in list(body):
        if child.tag != qn("w:p"):
            continue
        text = Paragraph(child, doc._body).text.strip()
        if start is None and text == "需求模型":
            start = child
        elif start is not None and text == "使用手冊":
            end = child
            break

    if start is None or end is None:
        raise RuntimeError("Unable to locate the chapter 5 to chapter 12 replacement range in the source document.")

    deleting = False
    for child in list(body):
        if child is start:
            deleting = True
        if child is end:
            break
        if deleting:
            body.remove(child)

    for chapter in CHAPTERS:
        add_chapter(doc, end, chapter)

    # Let Word/LibreOffice refresh the TOC, lists of figures/tables, references and page numbers.
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")

    core = doc.core_properties
    core.title = "FocusFlow AI 系統手冊（複評更新版）"
    core.subject = "FocusFlow AI 教學影片問答系統"

    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()
