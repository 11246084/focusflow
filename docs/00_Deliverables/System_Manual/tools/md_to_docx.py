"""Convert System Manual chapter Markdown to DOCX with the manual's typography.

pandoc converts Markdown (including LaTeX math in $...$ / $$...$$) to DOCX with
native Word equations; python-docx then applies the layout contract in
tools/artifact.md: A4 with 0.59-inch margins, DFKai-SB (標楷體) for Chinese and
Times New Roman for Latin text, headings 18/16/14 pt, body 14 pt justified with
a first-line indent, captions 14 pt centered, grid tables with a light-green
header row, and images scaled to the text width.

Usage:
    python md_to_docx.py <chapter.md> [<chapter.md> ...] --out <dir> [--pandoc <path>]

pandoc is looked up in this order: --pandoc, the PANDOC environment variable,
`pandoc` on PATH, and the pypandoc_binary package.
"""

from __future__ import annotations

import argparse
import os
import re
import shutil
import subprocess
import sys
import unicodedata
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt, RGBColor

CJK_FONT = "DFKai-SB"
LATIN_FONT = "Times New Roman"
BODY_PT = 14
TABLE_PT = 12
CODE_PT = 12
TEXT_WIDTH_IN = 7.0
MAX_IMAGE_W_IN = 6.75
MAX_IMAGE_H_IN = 8.25
HEADER_FILL = "D9EAD3"
CAPTION_RE = re.compile(r"^(圖|表|式)\s*\d+-\d+-\d+[a-z]?[\s　]")


def is_caption(text: str) -> bool:
    """Captions are short labels; body sentences that start with 表2-1-4 etc. are not."""
    return bool(CAPTION_RE.match(text)) and len(text) <= 60 and not re.search(r"[。，；：]", text)


def find_pandoc(explicit: str | None) -> str:
    candidates = [explicit, os.environ.get("PANDOC"), shutil.which("pandoc")]
    for candidate in candidates:
        if candidate and Path(candidate).exists():
            return candidate
    try:
        import pypandoc  # type: ignore

        return pypandoc.get_pandoc_path()
    except Exception as exc:  # pragma: no cover - environment dependent
        raise SystemExit("pandoc not found; install pandoc or pypandoc_binary, or pass --pandoc") from exc


def set_rfonts(r_pr) -> None:
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if r_fonts.get(qn(attr)) is not None:
            del r_fonts.attrib[qn(attr)]
    r_fonts.set(qn("w:ascii"), LATIN_FONT)
    r_fonts.set(qn("w:hAnsi"), LATIN_FONT)
    r_fonts.set(qn("w:cs"), LATIN_FONT)
    r_fonts.set(qn("w:eastAsia"), CJK_FONT)


def style_fonts(style, *, size=None, bold=None, color_black=True) -> None:
    r_pr = style.element.get_or_add_rPr()
    set_rfonts(r_pr)
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color_black:
        style.font.color.rgb = RGBColor(0, 0, 0)


def apply_styles(doc: Document) -> None:
    # Document defaults: every run without explicit fonts inherits these.
    styles_el = doc.styles.element
    doc_defaults = styles_el.find(qn("w:docDefaults"))
    if doc_defaults is not None:
        rpr_default = doc_defaults.find(qn("w:rPrDefault"))
        if rpr_default is not None and rpr_default.find(qn("w:rPr")) is not None:
            set_rfonts(rpr_default.find(qn("w:rPr")))

    for style in doc.styles:
        if style.type in (1, 2):  # paragraph and character styles
            try:
                style_fonts(style)
            except Exception:
                pass

    sizes = {
        "Normal": (BODY_PT, None),
        "Body Text": (BODY_PT, None),
        "First Paragraph": (BODY_PT, None),
        "Compact": (BODY_PT, None),
        "Caption": (BODY_PT, False),
        "Image Caption": (BODY_PT, False),
        "Table Caption": (BODY_PT, False),
        "Heading 1": (18, True),
        "Heading 2": (16, True),
        "Heading 3": (BODY_PT, True),
        "Heading 4": (BODY_PT, True),
        "Heading 5": (BODY_PT, True),
        "Source Code": (CODE_PT, None),
        # Inline code keeps the surrounding size; pandoc's default is smaller.
        "Verbatim Char": (BODY_PT, None),
    }
    for name, (size, bold) in sizes.items():
        try:
            style = doc.styles[name]
        except KeyError:
            continue
        style_fonts(style, size=size, bold=bold)
        style.font.italic = False

    h1 = doc.styles["Heading 1"].paragraph_format
    h1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.page_break_before = True
    for name in ("Heading 1", "Heading 2", "Heading 3", "Heading 4", "Heading 5"):
        try:
            pf = doc.styles[name].paragraph_format
        except KeyError:
            continue
        pf.keep_with_next = True
        pf.space_before = Pt(12)
        pf.space_after = Pt(6)


def set_page(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(section, side, Inches(0.59))


def has(element, tag: str) -> bool:
    return element.find(".//" + qn(tag)) is not None


def is_display_math(paragraph) -> bool:
    return paragraph._p.find(".//" + qn("m:oMathPara")) is not None


def is_image(paragraph) -> bool:
    return has(paragraph._p, "w:drawing")


def is_horizontal_rule(paragraph) -> bool:
    return not paragraph.text.strip() and (has(paragraph._p, "w:pict") or has(paragraph._p, "w:pBdr"))


def format_body_paragraphs(doc: Document) -> None:
    body_styles = {"Body Text", "First Paragraph", "Normal"}
    for paragraph in list(doc.paragraphs):
        if is_horizontal_rule(paragraph):
            paragraph._p.getparent().remove(paragraph._p)
            continue
        text = paragraph.text.strip()
        pf = paragraph.paragraph_format
        in_list = paragraph._p.pPr is not None and paragraph._p.pPr.numPr is not None
        if is_caption(text):
            paragraph.style = doc.styles["Caption"]
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.first_line_indent = Pt(0)
            pf.space_before = Pt(3)
            pf.space_after = Pt(9)
            if text.startswith("表"):
                pf.keep_with_next = True
            for run in paragraph.runs:
                run.font.size = Pt(BODY_PT)
                run.font.bold = False
            continue
        if is_image(paragraph):
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.first_line_indent = Pt(0)
            pf.keep_with_next = True
            continue
        if is_display_math(paragraph):
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.first_line_indent = Pt(0)
            pf.keep_with_next = True
            continue
        if paragraph.style.name in body_styles or (paragraph.style.name == "Compact" and in_list):
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf.line_spacing = 1.15
            pf.space_after = Pt(6)
            if not in_list:
                pf.first_line_indent = Pt(28)


def scale_images(doc: Document) -> None:
    max_w = Inches(MAX_IMAGE_W_IN)
    max_h = Inches(MAX_IMAGE_H_IN)
    for inline in doc.inline_shapes:
        w, h = inline.width, inline.height
        if not w or not h:
            continue
        # Fit every figure to the text box (scaling up or down) without distortion.
        ratio = min(max_w / w, max_h / h)
        inline.width = Emu(int(w * ratio))
        inline.height = Emu(int(h * ratio))


def visual_len(text: str) -> float:
    return sum(2.0 if unicodedata.east_asian_width(ch) in "WF" else 1.0 for ch in text)


def column_widths(table, total_in: float) -> list[float]:
    cols = len(table.columns)
    weights = []
    for c in range(cols):
        lengths = []
        for row in table.rows:
            try:
                lengths.append(visual_len(row.cells[c].text))
            except IndexError:
                lengths.append(0)
        header = lengths[0] if lengths else 4
        body = sorted(lengths[1:]) or [header]
        typical = body[int(len(body) * 0.75)] if len(body) > 1 else body[0]
        # The longest unbreakable token (file name, identifier, time range) must fit
        # on one line so Word does not split it mid-word.
        tokens = [
            visual_len(token)
            for row in table.rows[1:]
            if c < len(row.cells)
            for token in re.split(r"[\s，、。；：（）()]+", row.cells[c].text)
            if token
        ]
        longest_token = max(tokens) if tokens else 0
        # Give each column room for its header on one line and most of its cells on
        # one or two lines; very long cells wrap instead of starving other columns.
        weights.append(max(header + 2, min(typical, 44), longest_token + 1.5, 6.0))
    raw = [w / sum(weights) * total_in for w in weights]
    minimum = 0.8
    short = [i for i, w in enumerate(raw) if w < minimum]
    if short and len(short) < cols:
        deficit = sum(minimum - raw[i] for i in short)
        long_total = sum(raw[i] for i in range(cols) if i not in short)
        raw = [minimum if i in short else raw[i] - deficit * raw[i] / long_total for i in range(cols)]
    return raw


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    for old in tbl_pr.findall(qn("w:tblBorders")):
        tbl_pr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "000000")
        borders.append(node)
    tbl_pr.append(borders)


def format_tables(doc: Document) -> None:
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        set_table_borders(table)
        tbl_pr = table._tbl.tblPr
        for old in tbl_pr.findall(qn("w:tblW")):
            tbl_pr.remove(old)
        tbl_w = OxmlElement("w:tblW")
        tbl_w.set(qn("w:w"), str(int(TEXT_WIDTH_IN * 1440)))
        tbl_w.set(qn("w:type"), "dxa")
        tbl_pr.append(tbl_w)
        layout = tbl_pr.find(qn("w:tblLayout"))
        if layout is None:
            layout = OxmlElement("w:tblLayout")
            tbl_pr.append(layout)
        layout.set(qn("w:type"), "fixed")

        widths = column_widths(table, TEXT_WIDTH_IN)
        grid = table._tbl.tblGrid
        for i, col in enumerate(grid.findall(qn("w:gridCol"))):
            if i < len(widths):
                col.set(qn("w:w"), str(int(widths[i] * 1440)))

        short_table = len(table.rows) <= 15
        for r_idx, row in enumerate(table.rows):
            tr_pr = row._tr.get_or_add_trPr()
            tr_pr.append(OxmlElement("w:cantSplit"))
            if r_idx == 0:
                header = OxmlElement("w:tblHeader")
                header.set(qn("w:val"), "true")
                tr_pr.append(header)
            for c_idx, cell in enumerate(row.cells):
                if c_idx < len(widths):
                    cell.width = Inches(widths[c_idx])
                if r_idx == 0:
                    set_cell_shading(cell, HEADER_FILL)
                for paragraph in cell.paragraphs:
                    pf = paragraph.paragraph_format
                    # Keep short tables in one piece, and never leave a header row alone
                    # at the bottom of a page.
                    if r_idx == 0 or (short_table and r_idx < len(table.rows) - 1):
                        pf.keep_with_next = True
                    pf.first_line_indent = Pt(0)
                    pf.space_before = Pt(2)
                    pf.space_after = Pt(2)
                    pf.line_spacing = 1.0
                    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER if r_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
                    for run in paragraph.runs:
                        run.font.size = Pt(TABLE_PT)
                        if r_idx == 0:
                            run.font.bold = True


def convert(md_path: Path, out_dir: Path, pandoc: str) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / (md_path.stem + ".docx")
    subprocess.run(
        [
            pandoc,
            str(md_path.name),
            "-f",
            "gfm+tex_math_dollars",
            "-t",
            "docx",
            "--resource-path=.",
            "-o",
            str(out_path.resolve()),
        ],
        cwd=md_path.parent,
        check=True,
    )
    doc = Document(str(out_path))
    set_page(doc)
    apply_styles(doc)
    format_body_paragraphs(doc)
    scale_images(doc)
    format_tables(doc)
    doc.save(str(out_path))
    return out_path


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("chapters", nargs="+", type=Path)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--pandoc")
    args = parser.parse_args()
    pandoc = find_pandoc(args.pandoc)
    for chapter in args.chapters:
        print(convert(chapter.resolve(), args.out, pandoc))


if __name__ == "__main__":
    sys.exit(main())
